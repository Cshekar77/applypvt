import io
from datetime import datetime, timedelta, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from flask import make_response


# ── IST helper ────────────────────────────────────────────────────
IST = timezone(timedelta(hours=5, minutes=30))

def _to_ist_str(dt):
    if dt is None:
        return ''
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(IST).strftime('%d-%m-%Y %I:%M %p IST')


# ── Shared border ─────────────────────────────────────────────────
_thin        = Side(border_style='thin', color='CCCCCC')
_cell_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


# ═════════════════════════════════════════════════════════════════
#  EXPORT ALL APPLICATIONS
#  Columns match exactly: Timestamp, Email Address, Application Number,
#  Name of the Candidate, Mother's Name, Father's Name, Date of Birth,
#  Gender, Mobile Number of Parent, Mobile Number of Candidate,
#  E-mail Id, Address, Nationality, Religion, Category,
#  Candidate Belongs to HK Region, Kannada Medium, Rural Background,
#  Caste Certificate No., Parent's Annual Income, Income Certificate No.,
#  10th Standard Board, 10th Standard Percentage,
#  12th Standard Board, 12th Standard Stream, 12th Standard Combination,
#  12th Standard Max. Marks, 12th Standard Marks Scored,
#  12th Standard Percentage, Verified, Verified At
# ═════════════════════════════════════════════════════════════════
HEADERS = [
    'Timestamp',
    'Email Address',
    'Application Number',
    'Name of the Candidate',
    "Mother's Name",
    "Father's Name",
    'Date of Birth',
    'Gender',
    'Mobile Number of Parent',
    'Mobile Number of Candidate',
    'E-mail Id',
    'Address',
    'Nationality',
    'Religion',
    'Category',
    'Candidate Belongs to HK Region',
    'Kannada Medium',
    'Rural Background',
    'Caste Certificate No.',
    "Parent's Annual Income",
    'Income Certificate No.',
    '10th Standard Board',
    '10th Standard Percentage',
    '12th Standard Board',
    '12th Standard Stream',
    '12th Standard Combination',
    '12th Standard Max. Marks',
    '12th Standard Marks Scored',
    '12th Standard Percentage',
    'Verified',
    'Verified At',
]


def _bool(val):
    return 'Yes' if val else 'No'


def _row(app):
    return [
        _to_ist_str(app.submitted_at),                        # Timestamp
        app.email or '',                                       # Email Address
        app.application_number,                               # Application Number
        app.candidate_name or '',                             # Name of the Candidate
        app.mother_name or '',                                # Mother's Name
        app.father_name or '',                                # Father's Name
        app.dob or '',                                        # Date of Birth
        app.gender or '',                                     # Gender
        app.parent_mobile or '',                              # Mobile Number of Parent
        app.phone or '',                                      # Mobile Number of Candidate
        app.email or '',                                      # E-mail Id
        app.address or '',                                    # Address
        app.nationality or '',                                # Nationality
        app.religion or '',                                   # Religion
        app.category.name if app.category else '',            # Category
        _bool(app.hk_region),                                 # Candidate Belongs to HK Region
        _bool(app.kannada_medium),                            # Kannada Medium
        _bool(app.rural_background),                          # Rural Background
        app.caste_certificate_no or '',                       # Caste Certificate No.
        app.parent_annual_income or '',                       # Parent's Annual Income
        app.income_certificate_no or '',                      # Income Certificate No.
        app.board_10 or '',                                   # 10th Standard Board
        app.percent_10 or '',                                 # 10th Standard Percentage
        app.board_12 or '',                                   # 12th Standard Board
        app.stream_12 or '',                                  # 12th Standard Stream
        app.combination_12 or '',                             # 12th Standard Combination
        app.total_12 or '',                                   # 12th Standard Max. Marks
        app.obtained_12 or '',                                # 12th Standard Marks Scored
        app.percent_12 or '',                                 # 12th Standard Percentage
        'Yes' if app.is_verified else 'No',                   # Verified
        _to_ist_str(app.verified_at),                         # Verified At
    ]


def export_excel(applications, filename='students.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Applications'

    header_font  = Font(bold=True, color='FFFFFF', size=11)
    header_fill  = PatternFill('solid', fgColor='0F1E3C')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for col, h in enumerate(HEADERS, 1):
        cell           = ws.cell(row=1, column=col, value=h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align
        cell.border    = _cell_border
    ws.row_dimensions[1].height = 36

    for r, app in enumerate(applications, 2):
        for col, val in enumerate(_row(app), 1):
            cell           = ws.cell(row=r, column=col, value=val)
            cell.border    = _cell_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if r % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F5F0E8')

    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = make_response(buf.read())
    resp.headers['Content-Type']        = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition'] = f'attachment; filename={filename}'
    return resp


# ═════════════════════════════════════════════════════════════════
#  EXPORT VERIFIED STUDENTS ONLY
# ═════════════════════════════════════════════════════════════════
VERIFIED_HEADERS    = [
    'Rank',
    'Application Number',
    'Name of the Candidate',
    "Mother's Name",
    "Father's Name",
    'Date of Birth',
    'Gender',
    'Mobile Number of Candidate',
    'E-mail Id',
    'Category',
    '10th Standard Board',
    '10th Standard Percentage',
    '12th Standard Board',
    '12th Standard Stream',
    '12th Standard Combination',
    '12th Standard Max. Marks',
    '12th Standard Marks Scored',
    '12th Standard Percentage',
]
VERIFIED_COL_WIDTHS = [8, 18, 28, 24, 24, 16, 12, 22, 28, 18, 20, 14, 20, 18, 24, 14, 14, 14]


def export_verified_excel(applications, filename='verified_students.xlsx'):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Verified Students'

    num_cols = len(VERIFIED_HEADERS)
    end_col  = get_column_letter(num_cols)

    # Title row
    ws.merge_cells(f'A1:{end_col}1')
    title           = ws['A1']
    title.value     = 'BCA Admissions 2026–27  ·  Verified Students Rank List'
    title.font      = Font(bold=True, size=13, color='FFFFFF')
    title.fill      = PatternFill('solid', fgColor='0F1E3C')
    title.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Generated at
    ws.merge_cells(f'A2:{end_col}2')
    ts           = ws['A2']
    ts.value     = f'Generated on: {_to_ist_str(datetime.utcnow())}'
    ts.font      = Font(italic=True, size=9, color='555555')
    ts.fill      = PatternFill('solid', fgColor='E8EAF6')
    ts.alignment = Alignment(horizontal='right', vertical='center')
    ws.row_dimensions[2].height = 16

    # Column headers
    for col, h in enumerate(VERIFIED_HEADERS, 1):
        cell           = ws.cell(row=3, column=col, value=h)
        cell.font      = Font(bold=True, color='FFFFFF', size=11)
        cell.fill      = PatternFill('solid', fgColor='1A237E')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border    = _cell_border
    ws.row_dimensions[3].height = 36

    # Data rows
    for i, app in enumerate(applications, 1):
        row_num  = i + 3
        row_data = [
            app.rank if app.rank else '—',
            app.application_number,
            app.candidate_name or '—',
            app.mother_name or '—',
            app.father_name or '—',
            app.dob or '—',
            app.gender or '—',
            app.phone or '—',
            app.email or '—',
            app.category.name if app.category else '—',
            app.board_10 or '—',
            f"{app.percent_10:.2f}%" if app.percent_10 else '—',
            app.board_12 or '—',
            app.stream_12 or '—',
            app.combination_12 or '—',
            int(app.total_12) if app.total_12 else '—',
            int(app.obtained_12) if app.obtained_12 else '—',
            f"{app.percent_12:.2f}%" if app.percent_12 else '—',
        ]

        for col, val in enumerate(row_data, 1):
            cell           = ws.cell(row=row_num, column=col, value=val)
            cell.border    = _cell_border
            cell.alignment = Alignment(
                horizontal='center' if col in (1, 7, 10) else 'left',
                vertical='center',
                wrap_text=True,
            )
            if i % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F5F0E8')
        ws.row_dimensions[row_num].height = 22

    for col, width in enumerate(VERIFIED_COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = 'A4'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = make_response(buf.read())
    resp.headers['Content-Type']        = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition'] = f'attachment; filename={filename}'
    return resp


# ═════════════════════════════════════════════════════════════════
#  EXPORT PAYMENTS — One row per student, dynamic receipt columns
#  Receipt 1 No | Receipt 1 Amount | Receipt 1 Date |
#  Receipt 2 No | Receipt 2 Amount | Receipt 2 Date | ...
# ═════════════════════════════════════════════════════════════════
def export_payments_excel(students_data, filename='payment_overview.xlsx'):
    """
    students_data: list of dicts with keys:
        appl, allotment, fees, receipts
    One row per student. Receipt columns expand dynamically.
    """
    from datetime import timezone as tz

    def _ist(dt):
        if not dt:
            return ''
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone(IST).strftime('%d %b %Y, %I:%M %p')

    # Find max receipt count across all students
    max_receipts = max((len(d['receipts']) for d in students_data), default=0)
    max_receipts = max(max_receipts, 1)  # at least 1 receipt column group

    # Build headers
    fixed_headers = [
        'Rank',
        'Application Number',
        'Name of the Candidate',
        'Mobile Number of Candidate',
        'E-mail Id',
        'Category',
        'Quota',
        'Admit Category',
        'Admit Seat No',
        'Fee Structure',
        'Base Fee (₹)',
        'Additional Fee (₹)',
        'Total Fee (₹)',
        'Total Fees Set (₹)',
        'Total Paid (₹)',
        'Balance Remaining (₹)',
        'No. of Receipts',
    ]
    receipt_headers = []
    for i in range(1, max_receipts + 1):
        receipt_headers += [
            f'Receipt {i} No',
            f'Receipt {i} Amount (₹)',
            f'Receipt {i} Date',
        ]

    all_headers = fixed_headers + receipt_headers

    wb = Workbook()
    ws = wb.active
    ws.title = 'Payment Overview'

    # Header row styling
    header_fill      = PatternFill('solid', fgColor='1F3864')
    receipt_hdr_fill = PatternFill('solid', fgColor='0F4C75')
    header_font      = Font(color='FFFFFF', bold=True, size=11)
    thin_border      = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'),  bottom=Side(style='thin')
    )

    for col_idx, header in enumerate(all_headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = header_font
        cell.fill      = receipt_hdr_fill if header.startswith('Receipt') else header_fill
        cell.border    = thin_border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 36

    # Sort by rank
    students_data = sorted(students_data,
                           key=lambda x: (x['appl'].rank or 9999))

    alt_fill = PatternFill('solid', fgColor='EEF2FF')
    for row_idx, d in enumerate(students_data, 2):
        appl      = d['appl']
        a         = d['allotment']
        fees      = d['fees']
        receipts  = d['receipts']

        admit_cat_name = ''
        admit_seat_num = ''
        if a.admit_category_id and a.admit_category:
            admit_cat_name = a.admit_category.name
        if a.admit_seat_number:
            admit_seat_num = a.admit_seat_number

        fee_cat_name    = ''
        fee_base_amount = ''
        additional_fee  = ''
        total_fee       = ''
        if hasattr(appl, 'student_fee') and appl.student_fee:
            sf = appl.student_fee[0] if isinstance(appl.student_fee, list) else appl.student_fee
            if sf:
                if sf.fee_category:
                    fee_cat_name    = sf.fee_category.name
                    fee_base_amount = sf.fee_category.amount
                additional_fee = sf.additional_fee or 0
                base           = fee_base_amount if fee_base_amount != '' else 0
                total_fee      = base + additional_fee

        total_paid = sum(r.amount_paid for r in receipts) if receipts else 0
        balance    = (float(fees.total_fees) - total_paid) if fees and fees.total_fees else ''

        fixed_values = [
            appl.rank or '',
            appl.application_number,
            appl.candidate_name or '',
            appl.phone or '',
            appl.email or '',
            appl.category.name if appl.category else '',
            a.quota.title() if a.quota else '',
            admit_cat_name,
            admit_seat_num,
            fee_cat_name,
            fee_base_amount if fee_base_amount != '' else '',
            additional_fee if additional_fee != '' else '',
            total_fee if total_fee != '' else '',
            float(fees.total_fees) if fees and fees.total_fees else '',
            total_paid if receipts else 0,
            balance,
            len(receipts),
        ]

        # Receipt columns
        receipt_values = []
        for i in range(max_receipts):
            if i < len(receipts):
                r = receipts[i]
                receipt_ist = r.created_at
                if receipt_ist and receipt_ist.tzinfo is None:
                    receipt_ist = receipt_ist.replace(tzinfo=timezone.utc)
                receipt_date = receipt_ist.astimezone(IST).strftime('%d %b %Y') if receipt_ist else ''
                receipt_values += [r.receipt_number, r.amount_paid, receipt_date]
            else:
                receipt_values += ['', '', '']

        all_values = fixed_values + receipt_values
        fill = alt_fill if row_idx % 2 == 0 else None

        for col_idx, val in enumerate(all_values, 1):
            cell        = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
            if fill:
                cell.fill = fill

    # Auto-width
    for col_idx, header in enumerate(all_headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len    = len(str(header))
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 35)

    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = make_response(buf.read())
    resp.headers['Content-Type']        = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    resp.headers['Content-Disposition'] = f'attachment; filename={filename}'
    return resp